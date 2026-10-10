import hashlib
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
EXPORTER = ROOT / "scripts" / "export-fja-inventory.py"
VALIDATOR = ROOT / "scripts" / "validate-fja-inventory-export.py"

def make_item(external_id, title, body, details, source_url):
    html = f"<html><body><h1>{title}</h1><p>Original source capture for {external_id}</p></body></html>"
    html_bytes = html.encode("utf-8")
    html_hash = hashlib.sha256(html_bytes).hexdigest()
    details["sourceCapture"] = {
        "artifactPath": f"raw-html/{external_id}/{html_hash}.html",
        "htmlSha256": html_hash,
        "byteLength": len(html_bytes),
        "capturedAt": "2026-10-10T00:00:00Z",
        "normalizedTextSha256": hashlib.sha256(body.encode("utf-8")).hexdigest(),
    }
    return {
        "externalId": external_id,
        "sourceUrl": source_url,
        "title": title,
        "organizationName": "Example Organization",
        "publishedDate": "2026-10-01",
        "applicationStartDate": "2026-10-02",
        "applicationEndDate": "2026-10-30",
        "advertisementNumber": "EX/2026/01",
        "qualification": "Source-reported qualification",
        "vacancyCount": None,
        "detailStatus": "EXTRACTED",
        "sourceStatus": "UNKNOWN",
        "details": details,
        "rawText": body,
        "contentHash": hashlib.sha256((html_hash + body).encode("utf-8")).hexdigest(),
    }, html_bytes

def main():
    with tempfile.TemporaryDirectory(prefix="fja-export-test-") as temp:
        base = Path(temp)
        html_body = "Original normalized body content. " * 1300  # forces multi-cell chunking
        item1, html1 = make_item(
            "12345", "Example Multi-Post Recruitment", html_body,
            {
                "tableFields": {
                    "name of post": "Assistant\nClerk",
                    "total vacancies": "3",
                    "qualification": "Generic qualification; do not copy across posts",
                    "source-only field": "Preserve this verbatim",
                },
                "tables": [{
                    "tableIndex": 0,
                    "caption": "Post-wise details",
                    "rows": [
                        {"rowIndex": 0, "cells": ["Name of Post", "Vacancies", "Qualification", "Salary"]},
                        {"rowIndex": 1, "cells": ["Assistant", "2", "Graduate", "30000 per month"]},
                        {"rowIndex": 2, "cells": ["Clerk", "1", "12th pass", "25000 per month"]},
                    ],
                }],
                "allLinks": [{"label": "Official Notification PDF", "url": "https://example.gov.in/notice.pdf"}],
                "headings": [{"level": 1, "text": "Example Multi-Post Recruitment"}],
                "lists": [{"listType": "ul", "items": ["Bring ID proof", "Apply online"]}],
                "postNames": "Assistant\nClerk",
                "ageLimit": "18 to 35 years",
                "salary": "Source-reported salary summary",
                "location": "Delhi",
            },
            "https://www.freejobalert.com/articles/example-12345",
        )
        item2, html2 = make_item(
            "67890", "Example Single-Post Recruitment", "Single post body text",
            {
                "tableFields": {
                    "application start date": "02 Oct 2026",
                    "last date to apply": "30 Oct 2026",
                    "qualification": "Bachelor's degree",
                    "age limit": "18-30 years",
                    "application fee": "₹100",
                    "unmapped note": "Do not discard",
                },
                "tables": [],
                "allLinks": [{"label": "Apply Online", "url": "https://example.gov.in/apply"}],
                "headings": [{"level": 1, "text": "Example Single-Post Recruitment"}],
                "lists": [],
                "postNames": "Consultant",
                "ageLimit": "18-30 years",
                "salary": "₹40,000 per month",
                "location": "Gurugram",
            },
            "https://www.freejobalert.com/articles/example-67890",
        )
        # Keep the top-level convenience field consistent with the fixture's source table field.
        item2["qualification"] = "Bachelor's degree"
        (base / "raw-html" / "12345").mkdir(parents=True)
        (base / "raw-html" / "12345" / f"{item1['details']['sourceCapture']['htmlSha256']}.html").write_bytes(html1)
        (base / "raw-html" / "67890").mkdir(parents=True)
        (base / "raw-html" / "67890" / f"{item2['details']['sourceCapture']['htmlSha256']}.html").write_bytes(html2)
        payload = {
            "runId": "fja-test-run",
            "startedAt": "2026-10-10T00:00:00Z",
            "finishedAt": "2026-10-10T00:01:00Z",
            "results": [item1, item2],
        }
        (base / "fja-inventory.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        env = dict(os.environ, FJA_OUT_DIR=str(base))
        subprocess.run([sys.executable, str(EXPORTER)], check=True, env=env)
        subprocess.run([sys.executable, str(VALIDATOR)], check=True, env=env)

        workbook = load_workbook(base / "FJA_Jobs_Inventory.xlsx", read_only=True, data_only=True)
        assert workbook.sheetnames == ["Recruitments", "Posts"]
        recruitment_sheet = workbook["Recruitments"]
        recruitment_headers = [cell.value for cell in next(recruitment_sheet.iter_rows(min_row=1, max_row=1))]
        recruitment_rows = list(recruitment_sheet.iter_rows(min_row=2, values_only=True))
        body_columns = [i for i, header in enumerate(recruitment_headers) if str(header).startswith("source_body_text_part_")]
        row1 = next(row for row in recruitment_rows if row[recruitment_headers.index("source_external_id_internal")] == "12345")
        assert "".join(str(row1[i] or "") for i in body_columns) == html_body

        posts_sheet = workbook["Posts"]
        post_headers = [cell.value for cell in next(posts_sheet.iter_rows(min_row=1, max_row=1))]
        post_rows = list(posts_sheet.iter_rows(min_row=2, values_only=True))
        assert len(post_rows) == 3, f"Expected two table posts and one single post, got {len(post_rows)}"
        name_i = post_headers.index("source_post_name_raw")
        vacancy_i = post_headers.index("vacancy_count_normalized_candidate")
        qualification_i = post_headers.index("qualification_raw")
        key_i = post_headers.index("post_key")
        posts_by_name = {row[name_i]: row for row in post_rows}
        assert posts_by_name["Assistant"][vacancy_i] == 2
        assert posts_by_name["Assistant"][qualification_i] == "Graduate"
        assert posts_by_name["Clerk"][vacancy_i] == 1
        assert posts_by_name["Clerk"][qualification_i] == "12th pass"
        assert posts_by_name["Consultant"][qualification_i] == "Bachelor's degree"
        assert posts_by_name["Assistant"][key_i] == "12345-P01"
        assert posts_by_name["Clerk"][key_i] == "12345-P02"
        assert all(row[post_headers.index("official_verification_status")] == "PENDING_OFFICIAL_REVIEW" for row in post_rows)
        print("PASS: two-sheet export, source preservation, immutable HTML capture, and explicit post-row mappings")

if __name__ == "__main__":
    main()
