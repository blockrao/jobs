import hashlib
import json
import os
import re
from pathlib import Path
from openpyxl import load_workbook

base = Path(os.environ.get("FJA_OUT_DIR", "artifacts/fja"))
payload_path = base / "fja-inventory.json"
workbook_path = base / "FJA_Jobs_Inventory.xlsx"
if not payload_path.exists():
    raise SystemExit(f"Missing crawler payload: {payload_path}")
if not workbook_path.exists():
    raise SystemExit(f"Missing workbook: {workbook_path}")

payload = json.loads(payload_path.read_text(encoding="utf-8"))
items = payload.get("results", [])
candidate_path = base / "fja-post-candidates.json"
if not candidate_path.exists():
    raise SystemExit(f"Missing structured post candidate payload: {candidate_path}")
candidate_payload = json.loads(candidate_path.read_text(encoding="utf-8"))
if candidate_payload.get("runId") != payload.get("runId"):
    raise SystemExit("Recruitment and post candidate payload run IDs do not match")
expected_candidates = candidate_payload.get("postCandidates", [])
expected_by_key = {
    (str(candidate.get("externalId") or ""), str(candidate.get("sourcePostKey") or ""))
    : candidate
    for candidate in expected_output_candidates
}
if len(expected_by_key) != len(expected_candidates):
    raise SystemExit("Duplicate (externalId, sourcePostKey) in structured candidate payload")
wb = load_workbook(workbook_path, read_only=True, data_only=True)
expected_sheets = ["Recruitments", "Posts"]
if wb.sheetnames != expected_sheets:
    raise SystemExit(f"Expected exactly {expected_sheets}, found {wb.sheetnames}")

recruitments = wb["Recruitments"]
posts = wb["Posts"]
recruitment_headers = [cell.value for cell in next(recruitments.iter_rows(min_row=1, max_row=1))]
post_headers = [cell.value for cell in next(posts.iter_rows(min_row=1, max_row=1))]
if len(recruitment_headers) != len(set(recruitment_headers)):
    raise SystemExit("Duplicate header in Recruitments sheet")
if len(post_headers) != len(set(post_headers)):
    raise SystemExit("Duplicate header in Posts sheet")

recruitment_records = list(recruitments.iter_rows(min_row=2, values_only=True))
post_records = list(posts.iter_rows(min_row=2, values_only=True))
if len(recruitment_records) != len(items):
    raise SystemExit(f"Recruitment row mismatch: workbook={len(recruitment_records)} payload={len(items)}")
expected_output_candidates = [
    candidate for candidate in expected_candidates
    if str(candidate.get("externalId") or "") in {str(item.get("externalId") or "") for item in items}
]
if len(post_records) != len(expected_output_candidates):
    raise SystemExit(f"Post candidate row mismatch: workbook={len(post_records)} structured payload={len(expected_output_candidates)}")

recruitment_key_idx = recruitment_headers.index("recruitment_key")
source_id_idx = recruitment_headers.index("source_external_id_internal")
run_id_idx = recruitment_headers.index("run_id")
body_columns = [i for i, name in enumerate(recruitment_headers) if str(name).startswith("source_body_text_part_")]
details_columns = [i for i, name in enumerate(recruitment_headers) if str(name).startswith("source_details_json_part_")]
post_parent_idx = post_headers.index("recruitment_key")

parents = set()
recruitment_by_external_id = {}
for row in recruitment_records:
    key = row[recruitment_key_idx]
    external_id = str(row[source_id_idx] or "")
    if not key:
        raise SystemExit("Recruitment row missing recruitment_key")
    parents.add(key)
    recruitment_by_external_id[external_id] = row
    if row[run_id_idx] != payload.get("runId"):
        raise SystemExit(f"Run ID mismatch for {external_id}")

post_key_idx = post_headers.index("post_key")
post_external_id_idx = post_headers.index("source_external_id_internal")
post_name_idx = post_headers.index("source_post_name_raw")
post_candidate_count_idx = post_headers.index("vacancy_count_normalized_candidate")
post_salary_idx = post_headers.index("salary_raw")
post_qualification_idx = post_headers.index("qualification_raw")
post_extraction_status_idx = post_headers.index("extraction_status")
post_run_id_idx = post_headers.index("run_id")
seen_candidate_keys = set()
for row in post_records:
    parent = row[post_parent_idx]
    if parent not in parents:
        raise SystemExit(f"Post row has missing recruitment parent: {parent}")
    external_id = str(row[post_external_id_idx] or "")
    candidate_suffix = str(row[post_key_idx] or "")
    if not candidate_suffix.startswith(f"{external_id}-"):
        raise SystemExit(f"Post key is not scoped to its source article: {candidate_suffix}")
    source_post_key = candidate_suffix[len(external_id) + 1:]
    key = (external_id, source_post_key)
    if key in seen_candidate_keys:
        raise SystemExit(f"Duplicate post candidate in workbook: {key}")
    seen_candidate_keys.add(key)
    candidate = expected_by_key.get(key)
    if candidate is None:
        raise SystemExit(f"Workbook post row does not correspond to structured candidate: {key}")
    name = str(candidate.get("postNameRaw") or "").strip()
    if name and re.fullmatch(
        r"(?:pay|stipend|salary|application mode|job location|official website|total vacancies|vacancy|application fee|selection process|qualification|advertisement no\.?|organisation)",
        name,
        re.I,
    ):
        raise SystemExit(f"Recruitment-level label was exported as a post candidate: {name}")
    actual_name = str(row[post_name_idx] or "")
    if actual_name != (candidate.get("postNameRaw") or ""):
        raise SystemExit(f"Post title differs from structured candidate for {key}")
    expected_count = candidate.get("vacancyCountCandidate")
    actual_count = row[post_candidate_count_idx]
    if expected_count is not None and actual_count != expected_count:
        raise SystemExit(f"Post vacancy count was lost or changed for {key}: expected={expected_count}, actual={actual_count}")
    expected_salary = candidate.get("salaryRaw") or ""
    actual_salary = str(row[post_salary_idx] or "")
    if actual_salary != expected_salary:
        raise SystemExit(f"Post salary raw value was lost or changed for {key}")
    expected_qualification = candidate.get("qualificationRaw") or ""
    actual_qualification = str(row[post_qualification_idx] or "")
    if actual_qualification != expected_qualification:
        raise SystemExit(f"Post qualification was lost or changed for {key}")
    if row[post_extraction_status_idx] != candidate.get("extractionStatus"):
        raise SystemExit(f"Post extraction status mismatch for {key}")
    if row[post_run_id_idx] != payload.get("runId"):
        raise SystemExit(f"Post run ID mismatch for {key}")
if seen_candidate_keys != set(expected_by_key):
    missing = sorted(set(expected_by_key) - seen_candidate_keys)
    extra = sorted(seen_candidate_keys - set(expected_by_key))
    raise SystemExit(f"Workbook candidate key mismatch: missing={missing}, extra={extra}")

def joined_columns(row, indices):
    return "".join(str(row[index] or "") for index in indices)

for item in items:
    external_id = str(item.get("externalId") or "")
    row = recruitment_by_external_id.get(external_id)
    if row is None:
        raise SystemExit(f"Missing recruitment workbook row for source ID {external_id}")
    expected_body = str(item.get("rawText") or "")
    actual_body = joined_columns(row, body_columns)
    if actual_body != expected_body:
        raise SystemExit(f"Normalized body text was truncated or altered for source ID {external_id}")
    details = item.get("details") or {}
    expected_details = json.dumps(details, ensure_ascii=False, sort_keys=True, default=str)
    actual_details = joined_columns(row, details_columns)
    if actual_details != expected_details:
        raise SystemExit(f"Details JSON was truncated or altered for source ID {external_id}")
    capture = details.get("sourceCapture") or {}
    capture_rel_path = capture.get("artifactPath")
    expected_hash = capture.get("htmlSha256")
    if capture_rel_path and expected_hash:
        capture_path = base / capture_rel_path
        if not capture_path.exists():
            raise SystemExit(f"Missing original HTML capture for source ID {external_id}: {capture_path}")
        actual_hash = hashlib.sha256(capture_path.read_bytes()).hexdigest()
        if actual_hash != expected_hash:
            raise SystemExit(f"Original HTML hash mismatch for source ID {external_id}")
    elif item.get("detailStatus") != "FAILED":
        raise SystemExit(f"Successful article lacks immutable HTML capture metadata: {external_id}")

print(json.dumps({
    "status": "PASS",
    "runId": payload.get("runId"),
    "recruitmentRows": len(recruitment_records),
    "postCandidateRows": len(post_records),
    "structuredCandidateRowsMatched": len(seen_candidate_keys),
    "sheets": wb.sheetnames,
    "rawTextAndDetailsPreserved": True,
    "htmlCapturesVerified": True
}, ensure_ascii=False))
