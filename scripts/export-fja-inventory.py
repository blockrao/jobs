import csv
import json
import os
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = Path(os.environ.get("FJA_OUT_DIR", "artifacts/fja"))
PAYLOAD_PATH = BASE / "fja-inventory.json"
MAX_CELL_CHARS = 30000
payload = json.loads(PAYLOAD_PATH.read_text(encoding="utf-8"))
items = payload.get("results", [])
CANDIDATES_PATH = BASE / "fja-post-candidates.json"
if not CANDIDATES_PATH.exists():
    raise SystemExit(f"Missing structured post candidate payload: {CANDIDATES_PATH}")
candidate_payload = json.loads(CANDIDATES_PATH.read_text(encoding="utf-8"))
if candidate_payload.get("runId") != payload.get("runId"):
    raise SystemExit("Recruitment and post candidate payload run IDs do not match")
post_candidates = candidate_payload.get("postCandidates", [])
candidate_ids = {str(item.get("externalId") or "") for item in items}
unknown_candidate_ids = sorted({str(row.get("externalId") or "") for row in post_candidates} - candidate_ids)
if unknown_candidate_ids:
    raise SystemExit(f"Post candidates reference unknown recruitment IDs: {unknown_candidate_ids}")
candidates_by_external_id = {}
for candidate in post_candidates:
    candidates_by_external_id.setdefault(str(candidate.get("externalId") or ""), []).append(candidate)
BASE.mkdir(parents=True, exist_ok=True)

def as_json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str) if value is not None else ""

def as_text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return as_json(value)
    return str(value)

def chunks(value):
    value = as_text(value)
    return [value[i:i + MAX_CELL_CHARS] for i in range(0, len(value), MAX_CELL_CHARS)] if value else [""]

def chunk_headers(prefix, count):
    return [f"{prefix}_part_{i:02d}" for i in range(1, count + 1)]

def padded_chunks(value, count):
    parts = chunks(value) if value else []
    return (parts + [""] * count)[:count]

def first_field(fields, patterns):
    for key, value in (fields or {}).items():
        if value is not None and str(value).strip() and any(re.search(pattern, str(key), re.I) for pattern in patterns):
            return str(value).strip()
    return ""

def setup_sheet(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for index, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(index)].width = min(60, max(18, len(header) + 2))

def append_rows(ws, rows):
    for row in rows:
        ws.append(["" if value is None else value for value in row])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

def extract_links(details):
    # Prefer a full-link capture when the crawler provides it; retain legacy officialLinks otherwise.
    return details.get("allLinks") or details.get("officialLinks") or []

def classify_link_candidates(links):
    notification, application = [], []
    for link in links:
        label = str(link.get("label") or "")
        url = str(link.get("url") or "")
        if not url:
            continue
        if re.search(r"notification|advertisement|recruitment notice|detailed notice|official pdf", label, re.I):
            notification.append({"label": label, "url": url})
        if re.search(r"apply|application|registration|online form", label, re.I):
            application.append({"label": label, "url": url})
    return notification, application

def other_payload(item, details, fields, links):
    known_detail_keys = {
        "tableFields", "tables", "allLinks", "officialLinks", "headings", "lists",
        "sourceCapture", "postNames", "ageLimit", "applicationFee", "selectionProcess",
        "salary", "location", "rawText", "otherInfoRaw"
    }
    return {
        "unmapped_table_fields": {
            key: value for key, value in fields.items()
            if not re.search(
                r"organization|recruiting body|department|advertisement|notification|published|updated|application start|start date|starting date|last date|closing date|application end|total vacan|total post|number of vacan|no.? of post|qualification|eligibility|educational|application fee|exam fee|selection process|selection procedure|salary|pay scale|pay level|remuneration|job location|place of posting|location|name of post|post name",
                str(key), re.I
            )
        },
        "other_detail_fields": {key: value for key, value in details.items() if key not in known_detail_keys},
        "preservation_policy": "Raw source evidence only. No AI rewrite or official verification is performed by this exporter; source body, tables, links, headings and lists are retained in dedicated columns and the full details payload."
    }

recruitment_fixed_headers = [
    "recruitment_key", "import_action", "source_name_internal", "source_external_id_internal",
    "source_article_title_raw", "source_article_url_internal", "source_published_date_raw",
    "organization_name_raw", "organization_canonical_match_id", "official_notification_number_raw",
    "official_notification_url_candidate", "official_application_url_candidate",
    "recruitment_vacancy_total_raw", "recruitment_vacancy_total_normalized_candidate",
    "application_start_date_raw", "application_start_date_normalized_candidate",
    "application_deadline_date_raw", "application_deadline_date_normalized_candidate",
    "application_method_raw", "application_url_raw", "recruitment_location_raw",
    "recruitment_salary_summary_raw", "qualification_summary_raw", "fee_rules_raw_json",
    "selection_process_raw_json", "milestones_raw_json", "document_requirements_and_instructions_raw",
    "source_status_reported", "extraction_status", "official_verification_status",
    "canonical_promotion_gate", "promotion_gate_reason", "source_capture_ref_internal",
    "source_content_hash", "source_html_sha256", "run_id", "table_fields_json",
    "all_tables_json", "all_links_json", "headings_json", "lists_json", "unmapped_fields_json"
]
post_fixed_headers = [
    "post_key", "recruitment_key", "source_external_id_internal", "source_post_name_raw",
    "post_name_normalized_candidate", "source_post_code_raw", "vacancy_count_raw",
    "vacancy_count_normalized_candidate", "qualification_raw", "experience_raw",
    "minimum_age_years_candidate", "maximum_age_years_candidate", "age_limit_raw", "age_reference_date_raw",
    "age_relaxation_rules_raw", "salary_min_amount_candidate", "salary_max_amount_candidate",
    "salary_currency_candidate", "salary_period_candidate", "salary_raw", "pay_level_raw",
    "employment_type_raw", "tenure_raw", "location_raw", "duties_responsibilities_raw",
    "eligibility_conditions_raw", "post_milestones_json", "post_application_selection_json",
    "shared_source_fields_unresolved_json", "source_capture_ref_internal", "source_content_hash",
    "run_id", "post_structure_status", "extraction_status", "official_verification_status",
    "canonical_mapping_status", "review_note"
]

recruitment_base_rows = []
post_base_rows = []
recruitment_body_values = []
recruitment_details_values = []
recruitment_other_values = []
post_other_values = []

for item in items:
    details = item.get("details") or {}
    fields = details.get("tableFields") or {}
    links = extract_links(details)
    notification_links, application_links = classify_link_candidates(links)
    body = item.get("rawText") or details.get("rawText") or ""
    details_json = as_json(details)
    other_json = as_json(other_payload(item, details, fields, links))
    source_capture = details.get("sourceCapture") or {}
    external_id = str(item.get("externalId") or "")
    recruitment_key = f"FJA-{external_id}"
    start_raw = first_field(fields, [r"application start", r"start date", r"starting date"])
    end_raw = first_field(fields, [r"last date", r"closing date", r"application end", r"last date to apply"])
    fees = {key: value for key, value in fields.items() if re.search(r"fee", str(key), re.I)}
    selection = {key: value for key, value in fields.items() if re.search(r"selection", str(key), re.I)}
    milestones = {key: value for key, value in fields.items() if re.search(r"date|deadline|exam|interview|correction|age.*reckon|notification", str(key), re.I)}
    documents = {key: value for key, value in fields.items() if re.search(r"document|instruction|how to apply|application process", str(key), re.I)}
    notification_url = "; ".join(link["url"] for link in notification_links)
    application_url = "; ".join(link["url"] for link in application_links)
    gate = "BLOCKED_OFFICIAL_NOTIFICATION_REQUIRED" if not notification_links else "BLOCKED_PENDING_OFFICIAL_VERIFICATION"
    gate_reason = (
        "No likely notification link detected; retain inventory and find the issuing body's official notice during enrichment"
        if not notification_links else
        "Detected URL is only a candidate; validate issuing authority and facts before canonical promotion"
    )
    source_ref = source_capture.get("artifactPath") or f"source-observation:{external_id}:{item.get('contentHash','')}"
    recruitment_base = [
        recruitment_key, "REVIEW_ONLY", "freejobalert", external_id, item.get("title"),
        item.get("sourceUrl"), item.get("publishedDate"), item.get("organizationName"), "",
        item.get("advertisementNumber"), notification_url, application_url,
        first_field(fields, [r"total vacan", r"total post", r"number of vacan", r"no.? of post"]),
        item.get("vacancyCount"), start_raw, item.get("applicationStartDate"), end_raw,
        item.get("applicationEndDate"),
        first_field(fields, [r"application mode", r"how to apply", r"application process", r"apply mode"]),
        application_url,
        details.get("location") or first_field(fields, [r"job location", r"place of posting", r"location"]),
        details.get("salary") or first_field(fields, [r"salary", r"pay scale", r"pay level", r"remuneration", r"emolument"]),
        item.get("qualification") or first_field(fields, [r"qualification", r"eligibility", r"educational"]),
        as_json(fees), as_json(selection), as_json(milestones), as_json(documents),
        item.get("sourceStatus"), item.get("detailStatus"), "PENDING_OFFICIAL_REVIEW",
        gate, gate_reason, source_ref, item.get("contentHash"), source_capture.get("htmlSha256", ""),
        payload.get("runId"), as_json(fields), as_json(details.get("tables") or []), as_json(links),
        as_json(details.get("headings") or []), as_json(details.get("lists") or []),
        as_json(other_payload(item, details, fields, links).get("unmapped_table_fields"))
    ]
    recruitment_base_rows.append(recruitment_base)
    recruitment_body_values.append(chunks(body))
    recruitment_details_values.append(chunks(details_json))
    recruitment_other_values.append(chunks(other_json))

    candidates = candidates_by_external_id.get(external_id, [])
    for index, candidate in enumerate(candidates, 1):
        source_row = candidate.get("sourceTableRowRaw") or {}
        salary_normalization = candidate.get("salaryNormalization") or {}
        age_normalization = candidate.get("ageNormalization") or {}
        candidate_name = candidate.get("postNameRaw") or ""
        source_post_key = candidate.get("sourcePostKey") or f"unresolved-{index:02d}"
        post_key = f"{external_id}-{source_post_key}"
        scope_note = candidate.get("scopeNote") or ""
        extraction_status = candidate.get("extractionStatus") or "POST_DECOMPOSITION_NOT_EXTRACTED"
        post_other = {
            "candidate_name_raw": candidate_name,
            "source_post_key": source_post_key,
            "source_table_row_raw": source_row,
            "scope_note": scope_note,
            "source_post_candidate_hash": candidate.get("contentHash"),
            "salary_normalization_candidate": salary_normalization,
            "age_normalization_candidate": age_normalization,
            "preservation_policy": "Source candidate only; official verification is separate and canonical promotion is not authorized."
        }
        post_base_rows.append([
            post_key, recruitment_key, external_id, candidate_name, "", candidate.get("sourcePostCodeRaw") or "",
            candidate.get("vacancyCountRaw") or "", candidate.get("vacancyCountCandidate") if candidate.get("vacancyCountCandidate") is not None else "",
            candidate.get("qualificationRaw") or "", candidate.get("experienceRaw") or "",
            age_normalization.get("min") if age_normalization.get("min") is not None else "",
            age_normalization.get("max") if age_normalization.get("max") is not None else "",
            candidate.get("ageLimitRaw") or "", candidate.get("ageReferenceDateRaw") or "",
            candidate.get("ageRelaxationRulesRaw") or "",
            salary_normalization.get("min") if salary_normalization.get("min") is not None else "",
            salary_normalization.get("max") if salary_normalization.get("max") is not None else "",
            salary_normalization.get("currency") or "", salary_normalization.get("period") or "",
            candidate.get("salaryRaw") or "", candidate.get("payLevelRaw") or "",
            candidate.get("employmentTypeRaw") or "", candidate.get("tenureRaw") or "",
            candidate.get("locationRaw") or "", candidate.get("dutiesResponsibilitiesRaw") or "",
            candidate.get("eligibilityConditionsRaw") or "",
            as_json(candidate.get("milestonesRaw") or []), as_json(candidate.get("applicationSelectionRaw") or []),
            as_json(candidate.get("otherInfoRaw") or {}), source_ref, item.get("contentHash"),
            payload.get("runId"), extraction_status, item.get("detailStatus"),
            "PENDING_OFFICIAL_REVIEW", "SOURCE_TABLE_CANDIDATE" if source_row else "NAME_FIELD_CANDIDATE",
            scope_note
        ])
        post_other_values.append(chunks(as_json(post_other)))

max_body = max((len(value) for value in recruitment_body_values), default=1)
max_details = max((len(value) for value in recruitment_details_values), default=1)
max_recruitment_other = max((len(value) for value in recruitment_other_values), default=1)
max_post_other = max((len(value) for value in post_other_values), default=1)

recruitment_headers = (
    recruitment_fixed_headers + chunk_headers("source_body_text", max_body)
    + chunk_headers("source_details_json", max_details)
    + chunk_headers("other_info_raw", max_recruitment_other)
)
post_headers = post_fixed_headers + chunk_headers("post_other_info_raw", max_post_other)

recruitment_rows = []
for index, base_row in enumerate(recruitment_base_rows):
    recruitment_rows.append(
        base_row
        + padded_chunks("".join(recruitment_body_values[index]), max_body)
        + padded_chunks("".join(recruitment_details_values[index]), max_details)
        + padded_chunks("".join(recruitment_other_values[index]), max_recruitment_other)
    )
post_rows = [
    base_row + padded_chunks("".join(post_other_values[index]), max_post_other)
    for index, base_row in enumerate(post_base_rows)
]

# Ensure every output row has exactly as many cells as its header.
assert all(len(row) == len(recruitment_headers) for row in recruitment_rows), "Recruitments row/header width mismatch"
assert all(len(row) == len(post_headers) for row in post_rows), "Posts row/header width mismatch"

workbook = Workbook()
recruitments_ws = workbook.active
recruitments_ws.title = "Recruitments"
setup_sheet(recruitments_ws, recruitment_headers)
append_rows(recruitments_ws, recruitment_rows)
posts_ws = workbook.create_sheet("Posts")
setup_sheet(posts_ws, post_headers)
append_rows(posts_ws, post_rows)
workbook.properties.title = "JobOye Recruitment and Post Source Inventory"
workbook.properties.subject = "Two-table source inventory; not verified or approved for canonical publication"
workbook.properties.description = "Source-reported content preserved. Official verification and any AI rewriting are later, separate stages."
workbook.save(BASE / "FJA_Jobs_Inventory.xlsx")

def write_csv(path, headers, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)

write_csv(BASE / "Recruitments.csv", recruitment_headers, recruitment_rows)
write_csv(BASE / "Posts.csv", post_headers, post_rows)

summary = {
    "runId": payload.get("runId"),
    "recruitmentRows": len(recruitment_rows),
    "postCandidateRows": len(post_rows),
    "structuredPostCandidates": len(post_candidates),
    "unresolvedPostRows": sum(1 for row in post_rows if not row[0]),
    "workbookSheets": ["Recruitments", "Posts"],
    "sourceBodyTextChunkColumns": max_body,
    "sourceDetailsJsonChunkColumns": max_details,
    "recruitmentOtherInfoChunkColumns": max_recruitment_other,
    "postOtherInfoChunkColumns": max_post_other,
    "sourceReportedOnly": True,
    "officialVerificationComplete": False,
    "aiRewritingPerformed": False,
    "canonicalPromotionAuthorized": False,
    "limitations": [
        "Post decomposition is heuristic candidate data requiring review.",
        "Generic article-level facts are not copied onto each post for multi-post notices.",
        "Detected official/application URLs are candidates based on link labels, not verified URLs.",
        "The crawler must emit full links, structured table/list/heading content and immutable original HTML capture for maximum preservation."
    ]
}
(BASE / "export-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Workbook generated: {BASE / 'FJA_Jobs_Inventory.xlsx'} ({len(recruitment_rows)} recruitment rows; {len(post_rows)} post candidate rows; exactly two sheets)")
